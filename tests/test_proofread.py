"""The proofreading sheet's formulas agree with the site's numbers and its stored values."""
import pytest
from openpyxl import load_workbook

from bbgrid.config import WEB_DIR
from bbgrid.proofread import build

pycel = pytest.importorskip("pycel")


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    if not (WEB_DIR / "details.json").exists():
        pytest.skip("run `python -m bbgrid build` first")
    out, counts = build(out=tmp_path_factory.mktemp("sheet") / "proofread.xlsx")
    return out, counts


def test_formulas_match_stored_values(sheet):
    out, counts = sheet
    assert counts["players"] > 0 and counts["votes"] > 0
    formulas = load_workbook(out)
    stored = load_workbook(out, data_only=True)
    xl = pycel.ExcelCompiler(filename=str(out))
    checked = 0
    for title in ("Players", "Votes"):
        for row in formulas[title].iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    got = xl.evaluate(f"{title}!{cell.coordinate}")
                    want = stored[title][cell.coordinate].value
                    assert (got or "") == (want or ""), (title, cell.coordinate, cell.value, got, want)
                    checked += 1
    assert checked == counts["players"] * 8 + counts["votes"]
